"""Download URL extraction and small-file text readers."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urljoin

from bs4 import BeautifulSoup

from ntl_mcp.names import safe_folder

INLINE_LIMIT = 1_000_000
SHEET_ROW_CAP = 1000


def html_file_links(html: str, base_url: str) -> list[tuple[str, str]]:
    if not html:
        return []
    soup = BeautifulSoup(html, "html.parser")
    found: list[tuple[str, str]] = []
    seen: set[str] = set()
    for tag in soup.find_all("a", href=True):
        href = str(tag["href"]).strip()
        if not href or href.startswith("#") or href.lower().startswith("javascript:"):
            continue
        lowered = href.lower()
        useful = (
            "bbcswebdav" in lowered
            or "/xid-" in lowered
            or any(
                lowered.split("?", 1)[0].endswith(ext)
                for ext in (
                    ".pdf",
                    ".ppt",
                    ".pptx",
                    ".doc",
                    ".docx",
                    ".xls",
                    ".xlsx",
                    ".zip",
                    ".txt",
                    ".mp4",
                )
            )
        )
        if not useful:
            continue
        url = urljoin(base_url.rstrip("/") + "/", href)
        if url in seen:
            continue
        seen.add(url)
        label = tag.get_text(" ", strip=True) or unquote(url.split("?", 1)[0].rsplit("/", 1)[-1])
        found.append((url, safe_folder(label) or "download"))
    return found


def body_html(item: dict) -> str:
    body = item.get("body")
    if isinstance(body, str):
        return body
    if isinstance(body, dict):
        return str(body.get("rawText") or body.get("formattedText") or "")
    description = item.get("description")
    if isinstance(description, dict):
        return str(description.get("rawText") or description.get("formattedText") or "")
    if isinstance(description, str):
        return description
    return ""


def unique_path(directory: Path, filename: str) -> Path:
    name = safe_folder(filename)
    candidate = directory / name
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    n = 2
    while True:
        alt = directory / f"{stem} ({n}){suffix}"
        if not alt.exists():
            return alt
        n += 1


def read_small_file(data: bytes, filename: str, mime: str) -> str:
    if len(data) > INLINE_LIMIT:
        raise ValueError(
            f"File is {len(data)} bytes; over the {INLINE_LIMIT} inline cap. Use ntl_download_file."
        )
    name = filename.lower()
    mime = (mime or "").lower()
    if name.endswith(".pdf") or "pdf" in mime:
        return _pdf_text(data)
    if name.endswith(".docx") or "wordprocessingml" in mime:
        return _docx_text(data)
    if name.endswith(".xlsx") or "spreadsheetml" in mime:
        return _xlsx_text(data)
    if name.endswith(".pptx") or "presentationml" in mime:
        return _pptx_text(data)
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append(f"--- page {index} ---\n{(page.extract_text() or '').strip()}")
    text = "\n\n".join(pages).strip()
    return text or "[PDF had no extractable text]"


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text).strip() or "[empty docx]"


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in wb.worksheets:
        chunks.append(f"# {sheet.title}")
        for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if i > SHEET_ROW_CAP:
                chunks.append(f"... truncated after {SHEET_ROW_CAP} rows")
                break
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                chunks.append("\t".join(cells))
    return "\n".join(chunks).strip() or "[empty xlsx]"


def _pptx_text(data: bytes) -> str:
    from pptx import Presentation

    pres = Presentation(BytesIO(data))
    slides: list[str] = []
    for i, slide in enumerate(pres.slides, start=1):
        bits = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                bits.append(shape.text)
        slides.append(f"--- slide {i} ---\n" + "\n".join(bits))
    return "\n\n".join(slides).strip() or "[empty pptx]"
